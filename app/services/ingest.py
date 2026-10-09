"""Turn uploaded files or pasted text into searchable, citable chunks.

Keyword search works the moment a document is chunked; vectors are added in the
background by `embeddings.sync_embeddings` whenever the runtime is reachable.
"""

from __future__ import annotations

import io
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path

MAX_CHARS = 1400       # ~350 tokens: fits several sources into a small model's context
OVERLAP_CHARS = 200
SUPPORTED = {".pdf", ".docx", ".md", ".markdown", ".txt"}

_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


class IngestError(ValueError):
    pass


@dataclass(frozen=True)
class Page:
    number: int | None
    text: str


@dataclass(frozen=True)
class Chunk:
    page: int | None
    heading: str
    text: str


# -- parsing --------------------------------------------------------------------------

def parse_file(filename: str, data: bytes) -> list[Page]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise IngestError(f"Unsupported file type '{ext}'. Use PDF, DOCX, Markdown or text.")
    try:
        if ext == ".pdf":
            return _parse_pdf(data)
        if ext == ".docx":
            return _parse_docx(data)
        return [Page(None, data.decode("utf-8", errors="replace"))]
    except IngestError:
        raise
    except Exception as e:  # malformed files
        raise IngestError(f"Could not read {filename}: {e}") from e


def _parse_pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [Page(i + 1, (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages)]
    if not any(p.text for p in pages):
        raise IngestError("This PDF has no extractable text (it may be a scan).")
    return pages


def _parse_docx(data: bytes) -> list[Page]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = next((int(c) for c in style if c.isdigit()), 1)
            lines.append(f"{'#' * min(level, 6)} {text}")
        elif style == "title":
            lines.append(f"# {text}")
        else:
            lines.append(text)
        lines.append("")
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
        lines.append("")
    return [Page(None, "\n".join(lines))]


# -- chunking -------------------------------------------------------------------------

def chunk_pages(pages: list[Page]) -> list[Chunk]:
    """Split by heading and paragraph, keeping page + nearest heading for citations."""
    chunks: list[Chunk] = []
    heading = ""
    for page in pages:
        carry: list[str] = []  # short overlap from the previous chunk (same section)
        fresh: list[str] = []  # new paragraphs not yet emitted

        def flush(keep_overlap: bool) -> None:
            nonlocal carry, fresh
            if fresh:
                chunks.append(Chunk(page.number, heading, "\n\n".join(carry + fresh)))
                # carry a short tail so facts spanning a boundary stay retrievable
                tail = fresh[-1]
                carry = [tail] if keep_overlap and len(tail) <= OVERLAP_CHARS else []
            elif not keep_overlap:
                carry = []
            fresh = []

        for para in _paragraphs(page.text):
            m = _HEADING.match(para)
            if m:
                flush(keep_overlap=False)
                heading = m.group(2).strip()
                continue
            for piece in _split_long(para):
                size = sum(len(p) + 2 for p in carry + fresh)
                if fresh and size + len(piece) > MAX_CHARS:
                    flush(keep_overlap=True)
                fresh.append(piece)
        flush(keep_overlap=False)
    return chunks


def _paragraphs(text: str) -> list[str]:
    text = text.replace("\r\n", "\n")
    out: list[str] = []
    for block in re.split(r"\n\s*\n", text):
        lines = [ln.rstrip() for ln in block.strip().split("\n") if ln.strip()]
        if not lines:
            continue
        # A markdown heading glued to the following lines becomes its own paragraph.
        if _HEADING.match(lines[0]) and len(lines) > 1:
            out.append(lines[0])
            lines = lines[1:]
        out.append("\n".join(lines))
    return out


def _split_long(para: str) -> list[str]:
    if len(para) <= MAX_CHARS:
        return [para]
    sentences = re.split(r"(?<=[.!?])\s+", para)
    out, cur = [], ""
    for s in sentences:
        while len(s) > MAX_CHARS:  # pathological: no sentence breaks
            out.append(s[:MAX_CHARS])
            s = s[MAX_CHARS:]
        if len(cur) + len(s) + 1 > MAX_CHARS and cur:
            out.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        out.append(cur)
    return out


# -- persistence ----------------------------------------------------------------------

def store_document(
    db: sqlite3.Connection,
    *,
    title: str,
    filename: str,
    mime: str,
    data: bytes,
    owner_id: int | None,
    files_dir: Path,
) -> int:
    """Parse, chunk and store. Raises IngestError for unreadable input."""
    pages = parse_file(filename, data)
    chunks = chunk_pages(pages)
    if not chunks:
        raise IngestError("No text found in this document.")
    cur = db.execute(
        "INSERT INTO documents(title, filename, mime, size, owner_id, status, pages) "
        "VALUES (?, ?, ?, ?, ?, 'ready', ?)",
        (title.strip() or Path(filename).stem, filename, mime, len(data), owner_id,
         max((p.number or 0) for p in pages)),
    )
    doc_id = int(cur.lastrowid)
    db.executemany(
        "INSERT INTO chunks(document_id, ord, page, heading, text) VALUES (?, ?, ?, ?, ?)",
        [(doc_id, i, c.page, c.heading, c.text) for i, c in enumerate(chunks)],
    )
    files_dir.mkdir(parents=True, exist_ok=True)
    (files_dir / f"{doc_id}{Path(filename).suffix.lower()}").write_bytes(data)
    return doc_id


def delete_document(db: sqlite3.Connection, doc_id: int, files_dir: Path) -> bool:
    from ..db import vec_table

    row = db.execute("SELECT filename FROM documents WHERE id = ?", (doc_id,)).fetchone()
    if not row:
        return False
    chunk_ids = [r[0] for r in db.execute("SELECT id FROM chunks WHERE document_id = ?", (doc_id,))]
    for (space_id,) in db.execute("SELECT id FROM embedding_spaces").fetchall():
        db.executemany(f"DELETE FROM {vec_table(space_id)} WHERE rowid = ?", [(c,) for c in chunk_ids])
    db.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
    path = files_dir / f"{doc_id}{Path(row['filename']).suffix.lower()}"
    path.unlink(missing_ok=True)
    return True
