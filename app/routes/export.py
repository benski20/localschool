"""Word (.docx) export of any answer or draft. Nothing is stored.

    POST /api/export/docx   form fields: title, text, sources (JSON list of strings, optional)

A plain form post (not fetch + blob) so every browser, and the Mac app's web view, treats
the response as an ordinary download (Content-Disposition: attachment). The text is the
same plain text the Copy buttons produce; light markdown (# headings, - / 1. lists,
**bold**) is turned into Word formatting.
"""

from __future__ import annotations

import io
import json
import re
import sqlite3
from datetime import date
from urllib.parse import quote

from docx import Document
from docx.shared import Pt
from fastapi import APIRouter, Depends, Form
from fastapi.responses import Response

from .. import deps
from ..config import Settings

router = APIRouter(prefix="/api/export", tags=["export"])
MAX_TEXT = 200_000
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMBERED = re.compile(r"^\d+[.)]\s+")


def _runs(paragraph, text: str) -> None:
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build_docx(title: str, text: str, sources: list[str], school: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    header = doc.sections[0].header.paragraphs[0]
    header.text = " · ".join(p for p in (school, f"{date.today():%B} {date.today().day}, {date.today().year}") if p)
    if title.strip():
        doc.add_heading(title.strip(), level=1)
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.lstrip()
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")) + 1, 4)
            doc.add_heading(stripped.lstrip("#").strip(), level=level)
        elif stripped[:2] in ("- ", "* ", "• "):
            _runs(doc.add_paragraph(style="List Bullet"), stripped[2:].strip())
        elif _NUMBERED.match(stripped):
            _runs(doc.add_paragraph(style="List Number"), _NUMBERED.sub("", stripped, count=1))
        else:
            _runs(doc.add_paragraph(), line)
    if sources:
        doc.add_heading("Sources", level=2)
        for i, s in enumerate(sources, 1):
            doc.add_paragraph(f"[{i}] {s}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _filename(title: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", title).strip()[:80] or "Teacher's Desk"
    return f"{base}.docx"


@router.post("/docx")
def export_docx(
    title: str = Form(""),
    text: str = Form(..., max_length=MAX_TEXT),
    sources: str = Form("[]"),
    _: sqlite3.Row = Depends(deps.current_user),
    settings: Settings = Depends(deps.settings),
):
    try:
        src = [str(s)[:500] for s in json.loads(sources or "[]")][:100]
    except (ValueError, TypeError):
        src = []
    data = build_docx(title[:200], text, src, settings.school_name)
    name = _filename(title)
    ascii_name = name.encode("ascii", "ignore").decode() or "Teachers Desk.docx"
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}",
                 "Cache-Control": "no-store"},
    )
