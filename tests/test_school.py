"""School-readiness: Word export, privacy facts, activity log, Google/Microsoft sign-in."""

from __future__ import annotations

import io
import json

from docx import Document

from .conftest import login

JORDAN = "jordan.lee@shipley.demo"


def test_word_export_formats_and_needs_sign_in(client):
    form = {"title": "Ava — meeting notes", "text": "# Decisions\n- Daily **check-in** with Ms. Rivera\n1. Start Monday\nPlain line",
            "sources": json.dumps(["Handbook, p. 4"])}
    assert client.post("/api/export/docx", data=form).status_code == 401
    login(client, JORDAN)
    r = client.post("/api/export/docx", data=form)
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers["content-disposition"] and ".docx" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Heading 1", "Ava — meeting notes") in paras
    assert ("Heading 2", "Decisions") in paras
    assert ("List Bullet", "Daily check-in with Ms. Rivera") in paras
    assert ("List Number", "Start Monday") in paras
    assert any(t == "[1] Handbook, p. 4" for _, t in paras)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "check-in" in bold
    assert doc.sections[0].header.paragraphs[0].text  # school · date


def test_activity_log_records_who_and_whose_data_never_the_question(client, settings):
    from .conftest import sse

    client.post("/api/auth/login", json={"email": JORDAN, "password": "nope"})
    login(client, JORDAN)
    client.get("/api/roster"); client.get("/api/roster")  # second view within the hour isn't repeated
    question = "Which accommodations does Ava Thompson have for her secret fieldtrip?"
    sse(client.post("/api/chat", json={"question": question}))
    student = client.get("/api/roster").json()["students"][0]
    r = client.post("/api/tools/email/drafts", json={"student_id": student["id"], "topic": "great week"})
    assert r.status_code == 200, r.text
    assert client.get("/api/admin/activity").status_code == 403  # admins only

    login(client, "admin@test.local", "admin-pass-123")
    rows = client.get("/api/admin/activity").json()["activity"]
    acts = [(a["action"], a["who"], a["subject"]) for a in rows]
    assert ("sign_in_failed", "", JORDAN) in acts
    assert ("sign_in", "Jordan Lee", "") in acts
    assert sum(1 for a in acts if a[0] == "students_viewed") == 1
    assert any(a[0] == "students_used" and "Ava Thompson" in a[2] for a in acts)
    assert ("draft_for_student", "Jordan Lee", student["name"]) in acts
    flat = json.dumps(rows)
    assert "secret fieldtrip" not in flat  # questions are never logged

    assert any(a["subject"] == JORDAN for a in client.get("/api/admin/activity", params={"q": "jordan"}).json()["activity"])
    csv_text = client.get("/api/admin/activity.csv").text
    assert csv_text.startswith("when (UTC),who,what") and "Jordan Lee" in csv_text

    from app.db import connect
    from app.services import activity
    with connect(settings.db_path) as db:
        db.execute("UPDATE activity SET created_at = datetime('now', '-400 days') WHERE action = 'sign_in_failed'")
        assert activity.purge(db, 365) == 1


def test_trust_facts_match_the_install(client, settings):
    t = client.get("/api/trust").json()  # public: parents and boards read /privacy without signing in
    assert t["deployment"] == "school" and t["ai"] == "on_site" and t["activity_days"] == 365
    settings.demo_logins = True
    assert client.get("/api/trust").json()["deployment"] == "demo"
    settings.demo_logins, settings.local_mode = False, True
    assert client.get("/api/trust").json()["deployment"] == "personal"


# -- Sign in with Google / Microsoft, against a fake provider ---------------------------

import base64  # noqa: E402
import html  # noqa: E402
import re  # noqa: E402
import time  # noqa: E402
from urllib.parse import parse_qs, urlparse  # noqa: E402

import httpx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

IDP = {"google": {"auth": "https://idp.test/auth", "token": "https://idp.test/token", "issuer": "https://idp.test"}}


def _jwt(claims: dict) -> str:
    enc = lambda d: base64.urlsafe_b64encode(json.dumps(d).encode()).rstrip(b"=").decode()  # noqa: E731
    return f"{enc({'alg': 'RS256'})}.{enc(claims)}.sig"


def _setup_sso(client, settings, **extra):
    settings.sso_endpoints_override = json.dumps(IDP)
    login(client, "admin@test.local", "admin-pass-123")
    r = client.put("/api/admin/sso", json={"google_client_id": "cid-123", "google_client_secret": "shh-secret-xyz",
                                           "domains": "shipley.demo", **extra})
    assert r.status_code == 200 and r.json()["enabled"] == ["google"]
    assert "shh-secret" not in r.text and r.json()["redirect_uris"]["google"].endswith("/api/auth/sso/google/callback")
    client.post("/api/auth/logout")


def _provider(client, settings, mode="web", **claims):
    """Start a sign-in, then answer it as the provider would."""
    r = client.get(f"/api/auth/sso/google/start?mode={mode}", follow_redirects=False)
    assert r.status_code == 302, r.text
    q = parse_qs(urlparse(r.headers["location"]).query)
    assert q["client_id"] == ["cid-123"] and q["code_challenge_method"] == ["S256"] and q["hd"] == ["shipley.demo"]
    state = q["state"][0]
    from app.db import connect
    with connect(settings.db_path) as db:
        nonce = db.execute("SELECT nonce FROM sso_pending WHERE state = ?", (state,)).fetchone()["nonce"]
    base = {"iss": "https://idp.test", "aud": "cid-123", "exp": time.time() + 300, "nonce": nonce,
            "email": JORDAN, "email_verified": True, "name": "Jordan Lee"}
    base.update(claims)

    def token(req: httpx.Request) -> httpx.Response:
        body = parse_qs(req.content.decode())
        assert body["code_verifier"] and body["client_secret"] == ["shh-secret-xyz"]
        return httpx.Response(200, json={"id_token": _jwt(base), "access_token": "x"})

    client.app.state.sso_transport = httpx.MockTransport(token)
    return client.get(f"/api/auth/sso/google/callback?code=abc&state={state}", follow_redirects=False), state


def test_google_sign_in_in_the_browser(client, settings):
    assert "sso" not in client.get("/api/health").json()
    _setup_sso(client, settings)
    assert client.get("/api/health").json()["sso"] == ["google"]
    r, state = _provider(client, settings)
    assert r.status_code == 302 and r.headers["location"] == "/chat"
    assert client.get("/api/auth/me").json()["user"]["email"] == JORDAN
    # The same answer can't be used twice.
    again = client.get(f"/api/auth/sso/google/callback?code=abc&state={state}")
    assert "expired" in again.text


def test_google_sign_in_refuses_what_it_should(client, settings):
    _setup_sso(client, settings)
    for claims, says in [({"email": "someone@gmail.com"}, "Use your school account"),
                         ({"email": "new.teacher@shipley.demo"}, "doesn't have a Teacher's Desk account"),
                         ({"nonce": "forged"}, "couldn't be confirmed"),
                         ({"aud": "other-app"}, "different app"),
                         ({"iss": "https://evil.test"}, "unexpected place"),
                         ({"email_verified": False}, "isn't verified"),
                         ({"exp": time.time() - 3600}, "expired")]:
        c = TestClient(client.app)  # a fresh browser each time
        r, _ = _provider(c, settings, **claims)
        assert r.status_code == 200 and says in html.unescape(r.text), (claims, r.text[-300:])
        assert c.get("/api/auth/me").status_code == 401
    # Answer in a different browser than the one that started it.
    r = client.get("/api/auth/sso/google/start?mode=web", follow_redirects=False)
    state = parse_qs(urlparse(r.headers["location"]).query)["state"][0]
    other = TestClient(client.app)
    other.app.state.sso_transport = httpx.MockTransport(lambda req: httpx.Response(500))
    assert "different browser" in other.get(f"/api/auth/sso/google/callback?code=abc&state={state}").text


def test_auto_join_creates_a_teacher(client, settings):
    _setup_sso(client, settings, auto_join=True)
    r, _ = _provider(client, settings, email="new.teacher@shipley.demo", name="Riley New")
    assert r.headers["location"] == "/chat"
    me = client.get("/api/auth/me").json()["user"]
    assert me["name"] == "Riley New" and me["role"] == "teacher"


def test_mac_app_sign_in_hands_back_a_one_time_code(client, settings):
    _setup_sso(client, settings)
    r, _ = _provider(client, settings, mode="app")
    code = re.search(r"teachersdesk://signin\?code=([\w-]+)", r.text).group(1)
    app = TestClient(client.app)  # the Mac app's window, not the browser
    assert app.get(f"/api/auth/sso/redeem?code={code}", follow_redirects=False).headers["location"] == "/chat"
    assert app.get("/api/auth/me").json()["user"]["email"] == JORDAN
    reuse = TestClient(client.app)
    assert reuse.get(f"/api/auth/sso/redeem?code={code}", follow_redirects=False).headers["location"] == "/"

    # "Keep my data on this Mac" with Google: the code becomes this Mac's key.
    r, _ = _provider(client, settings, mode="desk")
    code = re.search(r"teachersdesk://desk\?code=([\w-]+)", r.text).group(1)
    d = client.post("/api/auth/device/sso", json={"code": code, "label": "Jordan's Mac"}).json()
    assert d["key"].startswith("ls_dev_") and d["email"] == JORDAN
    assert client.post("/api/auth/device/sso", json={"code": code}).status_code == 401
    # Only a personal desk hands out sign-in codes for a password.
    assert client.post("/api/auth/code", json={"email": JORDAN, "password": "demo-pass-123"}).status_code == 404
