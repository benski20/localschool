"""OllamaProvider against Ollama's real wire formats (mocked HTTP, no server)."""

from __future__ import annotations

import io
import json

import httpx
import pytest

from app.llm import ChatMessage, RuntimeError_
from app.llm.ollama import OllamaProvider
from app.llm.openai_compat import OpenAICompatProvider
from app.services.ingest import parse_file


def handler(request: httpx.Request) -> httpx.Response:
    assert request.headers["authorization"] == "Bearer tok"
    assert request.headers["modal-key"] == "k"
    path = request.url.path
    if path == "/api/chat":
        body = json.loads(request.content)
        assert body["stream"] is True and body["model"] in ("qwen3:8b", "plain")
        if body["model"] == "plain" and "think" in body:
            return httpx.Response(400, json={"error": '"plain" does not support thinking'})
        lines = [{"message": {"role": "assistant", "content": "", "thinking": "plan the answer"}, "done": False},
                 {"message": {"role": "assistant", "content": "Room "}, "done": False},
                 {"message": {"role": "assistant", "content": "114 [1]."}, "done": False},
                 {"message": {"role": "assistant", "content": ""}, "done": True}]
        return httpx.Response(200, text="\n".join(json.dumps(x) for x in lines) + "\n")
    if path == "/v1/chat/completions":
        body = json.loads(request.content)
        assert body["stream"] is True and body["model"] == "qwen3:8b"
        chunks = ["<think>", "plan the answer", "</think>", "Room ", "114 [1]."]
        lines = [f"data: {json.dumps({'choices': [{'delta': {'content': c}}]})}\n\n" for c in chunks]
        lines.append("data: [DONE]\n\n")
        return httpx.Response(200, text="".join(lines), headers={"content-type": "text/event-stream"})
    if path == "/v1/embeddings":
        inputs = json.loads(request.content)["input"]
        data = [{"index": i, "embedding": [float(i), 1.0]} for i in reversed(range(len(inputs)))]
        return httpx.Response(200, json={"data": data})
    if path == "/api/tags":
        return httpx.Response(200, json={"models": [
            {"name": "qwen3:8b", "size": 5_200_000_000,
             "details": {"parameter_size": "8.2B", "quantization_level": "Q4_K_M", "family": "qwen3"}},
            {"name": "nomic-embed-text:latest", "size": 274_000_000, "details": {}},
        ]})
    if path == "/api/ps":
        return httpx.Response(200, json={"models": [{"name": "qwen3:8b", "size_vram": 6_000_000_000}]})
    if path == "/api/version":
        return httpx.Response(200, json={"version": "0.12.0"})
    if path == "/api/pull":
        name = json.loads(request.content)["model"]
        if name == "nope":
            return httpx.Response(200, text=json.dumps({"error": "pull model manifest: file does not exist"}) + "\n")
        evts = [{"status": "pulling manifest"}, {"status": "downloading", "total": 10, "completed": 4},
                {"status": "downloading", "total": 10, "completed": 10}, {"status": "success"}]
        return httpx.Response(200, text="\n".join(json.dumps(e) for e in evts) + "\n")
    if path == "/api/delete":
        return httpx.Response(200)
    return httpx.Response(404, json={"error": "not found"})


@pytest.fixture
def provider():
    return OllamaProvider("https://demo--ollama.modal.run", token="tok", headers={"Modal-Key": "k"},
                          transport=httpx.MockTransport(handler))


async def test_chat_streams_visible_text_only(provider):
    out = [t async for t in provider.chat("qwen3:8b", [ChatMessage("user", "hi")])]
    assert "".join(out) == "Room 114 [1]."


async def test_chat_retries_without_think_for_models_that_reject_it(provider):
    out = [t async for t in provider.chat("plain", [ChatMessage("user", "hi")])]
    assert "".join(out) == "Room 114 [1]." and "plain" in provider._no_think_param


async def test_openai_compat_chat_strips_think_tags():
    p = OpenAICompatProvider("https://demo--ollama.modal.run", token="tok", headers={"Modal-Key": "k"},
                             transport=httpx.MockTransport(handler))
    out = [t async for t in p.chat("qwen3:8b", [ChatMessage("user", "hi")])]
    assert "".join(out) == "Room 114 [1]."


async def test_embed_orders_by_index(provider):
    assert await provider.embed("nomic-embed-text", ["a", "b", "c"]) == [[0.0, 1.0], [1.0, 1.0], [2.0, 1.0]]


async def test_models_health_pull_delete(provider):
    models = await provider.list_models()
    assert [m.name for m in models] == ["nomic-embed-text:latest", "qwen3:8b"]
    assert models[1].loaded and models[1].parameter_size == "8.2B"
    h = await provider.health()
    assert h.ok and h.version == "0.12.0" and h.vram_bytes == 6_000_000_000
    progress = [p async for p in provider.pull("qwen3:8b")]
    assert progress[-1].status == "success" and progress[2].completed == 10
    with pytest.raises(RuntimeError_, match="does not exist"):
        _ = [p async for p in provider.pull("nope")]
    await provider.delete("qwen3:8b")


async def test_auth_error_is_explained():
    p = OllamaProvider("http://x", transport=httpx.MockTransport(lambda r: httpx.Response(401, text="no")))
    with pytest.raises(RuntimeError_, match="credentials"):
        await p.list_models()
    h = await p.health()
    assert not h.ok and "credentials" in h.error


def test_parse_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Dress code", level=1)
    d.add_paragraph("Hats are not allowed indoors.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Period 1", "8:00"
    buf = io.BytesIO()
    d.save(buf)
    text = parse_file("policy.docx", buf.getvalue())[0].text
    assert "# Dress code" in text and "Hats are not allowed" in text and "Period 1 | 8:00" in text


# -- structured output + speech-to-text --------------------------------------------------

async def test_json_schema_is_sent_as_ollama_format_and_openai_response_format():
    seen = []

    def h(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        seen.append(body)
        if request.url.path == "/api/chat":
            return httpx.Response(200, text=json.dumps({"message": {"content": '{"a": 1}'}, "done": True}) + "\n")
        if "response_format" in body:  # a server without structured outputs
            return httpx.Response(400, json={"error": {"message": "response_format is not supported"}})
        return httpx.Response(200, text='data: {"choices": [{"delta": {"content": "{}"}}]}\n\ndata: [DONE]\n\n')

    schema = {"type": "object", "properties": {"a": {"type": "integer"}}}
    p = OllamaProvider("http://x", transport=httpx.MockTransport(h))
    assert "".join([t async for t in p.chat("qwen3:8b", [ChatMessage("user", "hi")], json_schema=schema)]) == '{"a": 1}'
    assert seen[-1]["format"] == schema and seen[-1]["think"] is False
    _ = [t async for t in p.chat("qwen3:8b", [ChatMessage("user", "hi")])]
    assert "format" not in seen[-1]

    o = OpenAICompatProvider("http://x", transport=httpx.MockTransport(h))
    assert "".join([t async for t in o.chat("m", [ChatMessage("user", "hi")], json_schema=schema)]) == "{}"
    assert seen[-2]["response_format"]["json_schema"]["schema"] == schema and "response_format" not in seen[-1]
    _ = [t async for t in o.chat("m", [ChatMessage("user", "hi")], json_schema=schema)]
    assert "response_format" not in seen[-1]  # remembered: no second 400


VERBOSE = {"task": "transcribe", "language": "en", "duration": 7.5, "text": "Hello there. We agreed.",
           "segments": [{"id": 0, "seek": 0, "start": 0.0, "end": 3.2, "text": " Hello there.", "tokens": [1],
                         "temperature": 0.0, "avg_logprob": -0.2, "compression_ratio": 1.1, "no_speech_prob": 0.01},
                        {"id": 1, "seek": 0, "start": 3.2, "end": 7.5, "text": " We agreed.  ", "tokens": [2],
                         "temperature": 0.0, "avg_logprob": -0.2, "compression_ratio": 1.1, "no_speech_prob": 0.01}]}


@pytest.mark.parametrize("cls", [OllamaProvider, OpenAICompatProvider])
async def test_transcribe_posts_multipart_with_auth(tmp_path, cls):
    audio = tmp_path / "meeting.webm"
    audio.write_bytes(b"\x1a\x45\xdf\xa3" + b"x" * 1000)
    seen = {}

    def h(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1/audio/transcriptions" and request.method == "POST"
        seen["auth"], seen["key"] = request.headers["authorization"], request.headers["modal-key"]
        seen["ctype"] = request.headers["content-type"]
        seen["body"] = request.read()
        # our sidecar's keep-alive whitespace comes first
        return httpx.Response(200, text="  \n " + json.dumps(VERBOSE), headers={"content-type": "application/json"})

    p = cls("https://rt.example", token="tok", headers={"Modal-Key": "k"}, transport=httpx.MockTransport(h))
    segs = await p.transcribe(audio, language="en", prompt="IEP meeting with Ava Thompson.")
    assert [(s.start, s.end, s.text) for s in segs] == [(0.0, 3.2, "Hello there."), (3.2, 7.5, "We agreed.")]
    assert seen["auth"] == "Bearer tok" and seen["key"] == "k" and seen["ctype"].startswith("multipart/form-data")
    body = seen["body"]
    for name, value in (("model", b"whisper"), ("response_format", b"verbose_json"), ("language", b"en"),
                        ("prompt", b"IEP meeting with Ava Thompson.")):
        assert f'name="{name}"'.encode() + b"\r\n\r\n" + value in body
    assert b'name="file"; filename="meeting.webm"' in body and b"Content-Type: audio/webm" in body
    assert b"x" * 1000 in body


async def test_transcribe_explains_missing_endpoint_and_errors(tmp_path):
    from app.llm import TRANSCRIBE_UNSUPPORTED

    audio = tmp_path / "a.wav"
    audio.write_bytes(b"RIFF....WAVE")
    for resp in (httpx.Response(404, text="404 page not found"), httpx.Response(405), httpx.Response(501)):
        p = OllamaProvider("http://x", transport=httpx.MockTransport(lambda r, resp=resp: resp))
        with pytest.raises(RuntimeError_) as e:
            await p.transcribe(audio)
        assert str(e.value) == TRANSCRIBE_UNSUPPORTED
    late = httpx.Response(200, text=" " + json.dumps({"error": {"message": "Invalid data found when processing input"}}))
    p = OllamaProvider("http://x", transport=httpx.MockTransport(lambda r: late))
    with pytest.raises(RuntimeError_, match="Transcription failed: Invalid data"):
        await p.transcribe(audio)
    plain = httpx.Response(200, json={"text": " Just text. ", "duration": 2.0})
    p = OllamaProvider("http://x", transport=httpx.MockTransport(lambda r: plain))
    assert [(s.start, s.end, s.text) for s in await p.transcribe(audio)] == [(0.0, 2.0, "Just text.")]
    p = OllamaProvider("http://x", transport=httpx.MockTransport(lambda r: httpx.Response(401)))
    with pytest.raises(RuntimeError_, match="credentials"):
        await p.transcribe(audio)
